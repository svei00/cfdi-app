"""
Pruebas de F2.9 "Exportar pestaña a CSV…" (ROADMAP.md D43-D45/§18) a nivel
MainWindow: habilitado/deshabilitado de la accion de menu, flujo completo
(default = columnas del visor, alterno = todas las columnas), y que
"Guardar Excel…"/"Guardar filtrado…" sigan comportandose EXACTAMENTE igual
que antes (criterio de aceptacion 7, §18.7) - esa regresion es la que este
archivo existe para prevenir.

Corre en modo offscreen (CI / entorno sin pantalla).
"""
import csv
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(REPO_ROOT, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

from PySide6.QtWidgets import QApplication  # noqa: E402

from cfdi_app import core, paths  # noqa: E402
from cfdi_app.ui.main_window import MainWindow  # noqa: E402
from cfdi_app.ui.state import AppState  # noqa: E402

FIXTURE_DIR = os.path.join(REPO_ROOT, "XML-Test")

_app = QApplication.instance() or QApplication([])


class _ExportCsvTestCase(unittest.TestCase):
    """Arranca la ventana YA en estado DONE con un resultado real cargado."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._orig = {
            "BASE_APP_DIR": paths.BASE_APP_DIR,
            "BOVEDA_XML_DIR": paths.BOVEDA_XML_DIR,
            "REPORTS_DIR": paths.REPORTS_DIR,
            "LOGS_DIR": paths.LOGS_DIR,
            "LAST_USED_DIR_FILE": paths.LAST_USED_DIR_FILE,
            "ASSETS_DIR": paths.ASSETS_DIR,
        }
        paths.BASE_APP_DIR = self._tmp.name
        paths.BOVEDA_XML_DIR = os.path.join(self._tmp.name, "BovedaCFDI")
        paths.REPORTS_DIR = os.path.join(self._tmp.name, "Reports")
        paths.LOGS_DIR = os.path.join(self._tmp.name, "Logs")
        paths.LAST_USED_DIR_FILE = os.path.join(paths.REPORTS_DIR, "last.txt")
        paths.ASSETS_DIR = os.path.join(self._tmp.name, "assets_vacios")

        self.window = MainWindow()
        self.result = core.process_path(FIXTURE_DIR)
        self.window._last_result = self.result
        self.window.results_view.set_result(self.result)
        self.window._apply_state(AppState.DONE)

    def tearDown(self):
        self.window.deleteLater()
        for key, value in self._orig.items():
            setattr(paths, key, value)
        self._tmp.cleanup()


class TestExportCsvActionEnablement(_ExportCsvTestCase):
    def test_enabled_when_done(self):
        self.assertTrue(self.window.export_csv_action.isEnabled())

    def test_disabled_outside_done_state(self):
        self.window._apply_state(AppState.PROCESSING)
        self.assertFalse(self.window.export_csv_action.isEnabled())

    def test_disabled_before_any_result(self):
        fresh = MainWindow()
        try:
            self.assertFalse(fresh.export_csv_action.isEnabled())
        finally:
            fresh.deleteLater()


class TestExportCsvFlow(_ExportCsvTestCase):
    def test_default_exports_curated_visor_columns_for_active_tab(self):
        # Criterio §18.7.3 (default) + §18.7.1 (nombre con pestaña).
        expected_path = os.path.join(self._tmp.name, "salida.csv")
        filtro_visor = "Columnas del visor (*.csv)"

        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(expected_path, filtro_visor)) as mock_dialog:
            self.window.on_export_csv_clicked()

        proposed_path = mock_dialog.call_args[0][2]
        self.assertIn("_Facturas", proposed_path)
        self.assertTrue(os.path.exists(expected_path))

        with open(expected_path, encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f))
        self.assertEqual(len(header), 23)  # columnas curadas del visor (D18)
        self.assertEqual(self.window._last_saved_path, expected_path)
        self.assertTrue(self.window.open_btn.isEnabled())

    def test_todas_las_columnas_filter_exports_full_column_set(self):
        # Criterio §18.7.3 (alterno).
        expected_path = os.path.join(self._tmp.name, "completo.csv")
        filtro_todas = "Todas las columnas (*.csv)"

        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(expected_path, filtro_todas)):
            self.window.on_export_csv_clicked()

        with open(expected_path, encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f))
        self.assertGreater(len(header), 40)  # las 70 columnas del Excel

    def test_respects_active_filter_on_the_active_tab(self):
        # Criterio §18.7.1: filtro activo -> solo esas filas, nombre con sufijo.
        invoice_panel = self.window.results_view._panels["invoice_data"]
        invoice_panel.search_box.setText("pintura")
        expected_path = os.path.join(self._tmp.name, "filtrado.csv")

        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(expected_path, "Columnas del visor (*.csv)")) as mock_dialog:
            self.window.on_export_csv_clicked()

        proposed_path = mock_dialog.call_args[0][2]
        self.assertIn("_Filtrado_pintura", proposed_path)

        with open(expected_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        filtradas = self.window.results_view.filtered_data()["invoice_data"]
        self.assertEqual(len(rows), len(filtradas))
        self.assertLess(len(rows), len(self.result.invoice_data))

    def test_exports_the_currently_active_tab_not_always_facturas(self):
        # Criterio §18.7.6.
        self.window.results_view.setCurrentIndex(1)  # Nomina
        expected_path = os.path.join(self._tmp.name, "nomina.csv")

        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(expected_path, "Columnas del visor (*.csv)")) as mock_dialog:
            self.window.on_export_csv_clicked()

        proposed_path = mock_dialog.call_args[0][2]
        self.assertIn("_Nomina", proposed_path)
        with open(expected_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), len(self.result.nomina_data))

    def test_cancelling_the_dialog_does_not_save_or_change_state(self):
        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=("", "")):
            self.window.on_export_csv_clicked()
        self.assertIsNone(self.window._last_saved_path)
        self.assertFalse(self.window.open_btn.isEnabled())

    def test_permission_error_shows_critical_dialog_not_crash(self):
        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(os.path.join(self._tmp.name, "bloqueado.csv"),
                                  "Columnas del visor (*.csv)")), \
             patch("cfdi_app.ui.main_window.core.export_tab_csv",
                   side_effect=PermissionError("archivo en uso")), \
             patch("cfdi_app.ui.main_window.QMessageBox.critical") as mock_critical:
            self.window.on_export_csv_clicked()
        mock_critical.assert_called_once()
        self.assertIsNone(self.window._last_saved_path)


class TestExcelExportsAreUnaffected(_ExportCsvTestCase):
    """Criterio de aceptacion 7 (§18.7): 'Guardar Excel…' y 'Guardar
    filtrado…' NO cambian con la introduccion del CSV."""

    def test_normal_excel_save_still_exports_all_three_sheets_worth_of_data(self):
        expected_path = os.path.join(self._tmp.name, "completo.xlsx")
        with patch("cfdi_app.ui.main_window.QFileDialog.getSaveFileName",
                    return_value=(expected_path, "")):
            self.window.on_save_clicked()

        import openpyxl
        wb = openpyxl.load_workbook(expected_path)
        self.assertIn("Invoices", wb.sheetnames)
        self.assertEqual(
            wb["Invoices"].max_row, 1 + len(self.result.invoice_data))

    def test_save_filtered_button_still_works_independently_of_csv_action(self):
        invoice_panel = self.window.results_view._panels["invoice_data"]
        invoice_panel.search_box.setText("pintura")
        self.assertTrue(self.window.save_filtered_btn.isEnabled())
        self.assertTrue(self.window.export_csv_action.isEnabled())


if __name__ == "__main__":
    unittest.main(verbosity=2)
