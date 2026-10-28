"""
Pruebas de F2.10 etapa 4 (docs/design-f2.10-desglose-conceptos.md §8, D56-D58):
combustible de monedero (complemento EstadoDeCuentaCombustible) en el
PT-Recibidas. Fixtures: tests/fixtures/edenred/ (anonimizados).
"""
import copy
import glob
import os
import sys
import tempfile
import unittest
from unittest import mock

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(REPO_ROOT, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

import openpyxl  # noqa: E402

from cfdi_app.parsers.xml_parser_40 import parse_cfdi_40_invoice  # noqa: E402
from cfdi_app.reports import pt_recibidas as pt  # noqa: E402
from cfdi_app.exporters import pt_exporter  # noqa: E402
from cfdi_app.exporters.excel_exporter import export_to_excel  # noqa: E402

EDENRED_DIR = os.path.join(REPO_ROOT, "tests", "fixtures", "edenred")
SUJETO = "GAGG010101000"
UUID_JUNIO = "3358E192-7DAB-45A3-AEBE-30EADA074BDC"
UUID_AGOSTO = "46613D69-8126-672B-547F-3F98CD5132A6"

_CITAS_COMPLETAS = [(i, c, "cita verificada de prueba")
                    for i, c, _ in pt.FUNDAMENTOS_ECC]


def _registros():
    return [parse_cfdi_40_invoice(f)
            for f in sorted(glob.glob(os.path.join(EDENRED_DIR, "*.xml")))]


def _por_uuid(registros, uuid):
    return next(r for r in registros if r["UUID"] == uuid)


def _mes(resumen, clave):
    return next(f for f in resumen if f["mes"] == clave)


class TestCifrasDelMonedero(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.registros = _registros()
        cls.resumen = pt.construir_resumen_mensual(cls.registros, SUJETO)
        cls.detalle = pt.construir_detalle(cls.registros, SUJETO)

    def test_junio_y_agosto_traen_las_cifras_del_complemento(self):
        junio = _mes(self.resumen, "2026-06")
        self.assertAlmostEqual(junio["combustible_monedero"], 28185.69, 2)
        self.assertAlmostEqual(junio["iva_combustible_monedero"], 3791.92, 2)
        agosto = _mes(self.resumen, "2026-08")
        self.assertAlmostEqual(agosto["combustible_monedero"], 40447.09, 2)
        self.assertAlmostEqual(agosto["iva_combustible_monedero"], 5441.49, 2)

    def test_septiembre_sin_estado_de_cuenta_queda_en_cero(self):
        sept = _mes(self.resumen, "2026-09")
        self.assertEqual(sept["combustible_monedero"], 0.0)
        self.assertEqual(sept["iva_monedero_acreditable"], 0.0)

    def test_la_comision_no_cambia_de_clasificacion(self):
        # D58: lo que ya existia vale lo mismo. El IVA acreditable estimado de
        # junio sigue siendo SOLO el de la comision (148.80 + lo ya existente).
        junio = _mes(self.resumen, "2026-06")
        self.assertAlmostEqual(junio["iva_acreditable_estimado"], 148.80, 2)
        self.assertAlmostEqual(junio["total_neto"], 1078.80, 2)

    def test_detalle_trae_cargas_litros_e_importes(self):
        fila = next(d for d in self.detalle if d["uuid"] == UUID_JUNIO)
        self.assertEqual(fila["ecc_cargas"], 33)
        self.assertAlmostEqual(fila["ecc_litros"], 1127.97, 2)
        self.assertAlmostEqual(fila["ecc_importe_mxn"], 24393.77, 2)
        self.assertAlmostEqual(fila["ecc_iva_mxn"], 3791.92, 2)
        self.assertAlmostEqual(fila["ecc_total_mxn"], 28185.69, 2)
        self.assertEqual(fila["factura"], "TCGE-000004231229")

    def test_factura_sin_complemento_tiene_todo_en_cero(self):
        fila = next(d for d in self.detalle if d["ecc_cargas"] == 0)
        self.assertEqual(fila["ecc_total_mxn"], 0.0)
        self.assertEqual(fila["ecc_estatus"], "")
        self.assertEqual(fila["ecc_nivel"], "")
        self.assertFalse(fila["ecc_timbrado_otro_mes"])

    def test_el_timbrado_de_ambos_estados_cae_en_el_mes_siguiente(self):
        # Hallazgo real: emitidos el ultimo dia del mes y timbrados al dia
        # siguiente. La regla 3.3.1.7 dice "hasta que cuente con el CFDI".
        for uuid, emision, timbrado in (
                (UUID_JUNIO, (2026, 6), (2026, 7)),
                (UUID_AGOSTO, (2026, 8), (2026, 9))):
            fila = next(d for d in self.detalle if d["uuid"] == uuid)
            self.assertEqual(fila["mes"], emision)
            self.assertEqual(fila["mes_timbrado"], timbrado)
            self.assertTrue(fila["ecc_timbrado_otro_mes"])


class TestCompuertaFiscal(unittest.TestCase):
    """D56: las cifras se muestran, pero solo suman al IVA acreditable con
    F1-F4 verificados."""

    @classmethod
    def setUpClass(cls):
        cls.registros = _registros()

    def test_f1_y_f3_citan_la_regla_3_3_1_7_y_f2_f4_siguen_pendientes(self):
        # ACTUALIZAR esta prueba cuando se verifiquen F2 (LISR 27-III parr. 2)
        # y F4 (permiso de la Ley de Hidrocarburos): es la señal de revision.
        citas = {i: c for i, _, c in pt.FUNDAMENTOS_ECC}
        self.assertIn("3.3.1.7", citas["F1"])
        self.assertIn("monederos electrónicos autorizados por el SAT", citas["F1"])
        self.assertIn("hasta por el monto que ampare el citado complemento",
                      citas["F3"])
        self.assertIsNone(citas["F2"])
        self.assertIsNone(citas["F4"])

    def test_fundamentado_depende_de_las_cuatro_citas(self):
        self.assertEqual(pt.ecc_fundamentado(),
                         all(c is not None for _, _, c in pt.FUNDAMENTOS_ECC))
        self.assertFalse(pt.ecc_fundamentado())  # hoy faltan F2 y F4

    def test_compuerta_abierta_muestra_pero_no_suma(self):
        resumen = pt.construir_resumen_mensual(self.registros, SUJETO)
        junio = _mes(resumen, "2026-06")
        self.assertAlmostEqual(junio["iva_combustible_monedero"], 3791.92, 2)
        self.assertEqual(junio["iva_monedero_acreditable"], 0.0)
        fila = next(d for d in pt.construir_detalle(self.registros, SUJETO)
                    if d["uuid"] == UUID_JUNIO)
        self.assertEqual(fila["ecc_estatus"], pt.ESTATUS_ECC_PENDIENTE)
        self.assertEqual(fila["ecc_nivel"], pt.NIVEL_AMBAR)
        self.assertEqual(fila["ecc_iva_acreditable_mxn"], 0.0)

    def test_compuerta_cerrada_suma_y_pone_verde(self):
        # Se parchea SOLO en la prueba; el codigo de produccion no se toca.
        with mock.patch.object(pt, "FUNDAMENTOS_ECC", _CITAS_COMPLETAS):
            self.assertTrue(pt.ecc_fundamentado())
            resumen = pt.construir_resumen_mensual(self.registros, SUJETO)
            fila = next(d for d in pt.construir_detalle(self.registros, SUJETO)
                        if d["uuid"] == UUID_JUNIO)
        junio = _mes(resumen, "2026-06")
        self.assertAlmostEqual(junio["iva_monedero_acreditable"], 3791.92, 2)
        self.assertEqual(fila["ecc_estatus"], pt.ESTATUS_ECC_DEDUCIBLE)
        self.assertEqual(fila["ecc_nivel"], pt.NIVEL_VERDE)
        # La columna "IVA acreditable estimado" de siempre NO absorbe el ECC.
        self.assertAlmostEqual(junio["iva_acreditable_estimado"], 148.80, 2)

    def test_comision_a_xaxx_pone_rojas_las_cargas(self):
        regs = copy.deepcopy(self.registros)
        _por_uuid(regs, UUID_JUNIO)["RFC Receptor"] = "XAXX010101000"
        fila = next(d for d in pt.construir_detalle(regs, SUJETO)
                    if d["uuid"] == UUID_JUNIO)
        self.assertEqual(fila["ecc_estatus"], pt.ESTATUS_NO_DED_GENERICO)
        self.assertEqual(fila["ecc_nivel"], pt.NIVEL_ROJO)
        self.assertEqual(fila["ecc_iva_acreditable_mxn"], 0.0)

    def test_comision_ppd_sin_rep_pone_ambar_provisionado(self):
        regs = copy.deepcopy(self.registros)
        _por_uuid(regs, UUID_JUNIO)["Metodo de Pago"] = "PPD - Pago en parcialidades"
        with mock.patch.object(pt, "FUNDAMENTOS_ECC", _CITAS_COMPLETAS):
            fila = next(d for d in pt.construir_detalle(regs, SUJETO)
                        if d["uuid"] == UUID_JUNIO)
        self.assertEqual(fila["ecc_estatus"], pt.ESTATUS_PROVISIONADO)
        self.assertEqual(fila["ecc_nivel"], pt.NIVEL_AMBAR)
        self.assertEqual(fila["ecc_iva_acreditable_mxn"], 0.0)


class TestPosibleDuplicado(unittest.TestCase):
    """§8.3: factura de una estacion que repite una carga del monedero."""

    def _lote(self, total, fecha_emision, rfc="OMC100818GZ7"):
        regs = _registros()
        regs.append({
            "CFDI_Type": "Invoice", "Tipo": "Factura",
            "UUID": "CCCCCCCC-0000-0000-0000-000000000001",
            "RFC Emisor": rfc, "Nombre Emisor": "GASOLINERA",
            "RFC Receptor": SUJETO, "Metodo de Pago": "PUE - x",
            "FormaDePago": "03 - x", "Fecha Emision": fecha_emision,
            "SubTotal": round(total / 1.16, 2), "IVA 16%": 0.0,
            "Total": total, "Tipo De Cambio": 1.0, "ConceptosDetalle": [],
        })
        det = pt.construir_detalle(regs, SUJETO)
        return next(d for d in det if d["uuid"].startswith("CCCCCCCC"))

    def test_mismo_importe_y_dias_despues_se_marca_con_el_folio(self):
        # Carga 543884498: 2026-06-24, 1,251.39 + 194.49 = 1,445.88.
        fila = self._lote(1445.88, "25/06/2026")
        self.assertEqual(fila["posible_duplicado_monedero"], "TCGE-000004231229")

    def test_importe_distinto_no_se_marca(self):
        self.assertEqual(
            self._lote(1445.00, "25/06/2026")["posible_duplicado_monedero"], "")

    def test_fuera_de_la_ventana_no_se_marca(self):
        self.assertEqual(
            self._lote(1445.88, "10/08/2026")["posible_duplicado_monedero"], "")

    def test_emitida_antes_de_la_carga_no_se_marca(self):
        self.assertEqual(
            self._lote(1445.88, "20/06/2026")["posible_duplicado_monedero"], "")

    def test_otra_estacion_no_se_marca(self):
        self.assertEqual(
            self._lote(1445.88, "25/06/2026", rfc="ZZZ010101ZZZ")[
                "posible_duplicado_monedero"], "")

    def test_el_estado_de_cuenta_no_se_marca_a_si_mismo(self):
        det = pt.construir_detalle(_registros(), SUJETO)
        self.assertTrue(all(d["posible_duplicado_monedero"] == "" for d in det))


class TestExcelDelPt(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.registros = _registros()
        ruta = os.path.join(tempfile.mkdtemp(), "pt.xlsx")
        export_to_excel(cls.registros, [], [], ruta, incluir_pt=True)
        cls.wb = openpyxl.load_workbook(ruta)

    def test_hojas_del_pt_existen(self):
        for nombre in (pt_exporter.HOJA_RESUMEN, pt_exporter.HOJA_DETALLE,
                       pt_exporter.HOJA_FUNDAMENTOS):
            self.assertIn(nombre, self.wb.sheetnames)

    def test_resumen_trae_columnas_nuevas_y_formula_viva(self):
        ws = self.wb[pt_exporter.HOJA_RESUMEN]
        enc = [c.value for c in ws[1]]
        self.assertIn("Combustible monedero", enc)
        self.assertIn("IVA combustible monedero", enc)
        self.assertTrue(any(str(e).startswith("IVA monedero acreditable")
                            and "pendiente de verificar fundamento" in str(e)
                            for e in enc))
        col_formula = enc.index("IVA acreditable estimado + monedero") + 1
        col_iva = enc.index("IVA acreditable estimado") + 1
        col_mon = next(i for i, e in enumerate(enc, start=1)
                       if str(e).startswith("IVA monedero acreditable"))
        celda = ws.cell(row=2, column=col_formula).value
        letra = openpyxl.utils.get_column_letter
        self.assertEqual(celda, f"={letra(col_iva)}2+{letra(col_mon)}2")
        # El monto del monedero esta en junio (fila 2) como valor, no formula.
        self.assertAlmostEqual(
            ws.cell(row=2, column=enc.index("Combustible monedero") + 1).value,
            28185.69, 2)

    def test_las_columnas_de_siempre_no_se_movieron(self):
        ws = self.wb[pt_exporter.HOJA_RESUMEN]
        enc = [c.value for c in ws[1]]
        self.assertEqual(enc[0], "Mes")
        self.assertEqual(enc[12], "IVA acreditable estimado")
        det = [c.value for c in self.wb[pt_exporter.HOJA_DETALLE][1]]
        self.assertEqual(det[14], "Conciliado banco")
        self.assertEqual(det[19], "Estatus ECC")

    def test_detalle_marca_estatus_ecc_y_timbrado(self):
        ws = self.wb[pt_exporter.HOJA_DETALLE]
        filas = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        junio = next(f for f in filas if f[1] == UUID_JUNIO)
        self.assertEqual(junio[15], 33)
        self.assertEqual(junio[19], pt.ESTATUS_ECC_PENDIENTE)
        self.assertEqual(junio[20], "2026-07")

    def test_fundamentos_cita_f1_f3_y_marca_pendientes_f2_f4(self):
        ws = self.wb[pt_exporter.HOJA_FUNDAMENTOS]
        texto = {r[0].value: r[2].value for r in ws.iter_rows()
                 if r[0].value in ("F1", "F2", "F3", "F4")}
        self.assertIn("3.3.1.7", texto["F1"])
        self.assertIn("3.3.1.7", texto["F3"])
        self.assertIn("PENDIENTE DE VERIFICAR", texto["F2"])
        self.assertIn("PENDIENTE DE VERIFICAR", texto["F4"])


class TestSinMonedero(unittest.TestCase):
    def test_lote_sin_ecc_deja_las_columnas_nuevas_en_cero(self):
        sin = [r for r in _registros() if not r["EstadoCuentaCombustible"]]
        resumen = pt.construir_resumen_mensual(sin, SUJETO)
        self.assertGreater(len(resumen), 0)
        for f in resumen:
            self.assertEqual(f["combustible_monedero"], 0.0)
            self.assertEqual(f["iva_combustible_monedero"], 0.0)
            self.assertEqual(f["iva_monedero_acreditable"], 0.0)


if __name__ == "__main__":
    unittest.main()
