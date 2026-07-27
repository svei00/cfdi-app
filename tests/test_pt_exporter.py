"""
Pruebas del exporter del Papel de Trabajo (F2.8 etapa 1) a nivel de LIBRO de
Excel real: que las hojas se creen, que el gate las incluya/excluya, que la
columna "Conciliado banco" lleve formula VIVA, y que el resumen CUADRE contra
los datos de origen.

Cubre criterios de aceptacion §17.9: 1 (cuadre), 6 (Conciliado banco vivo),
7 (gate Lite = Excel sin hojas PT, regresion).
"""
import os
import sys
import tempfile
import unittest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(REPO_ROOT, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

import openpyxl  # noqa: E402

from cfdi_app.exporters.excel_exporter import export_to_excel  # noqa: E402
from cfdi_app.exporters import pt_exporter  # noqa: E402


def _factura(uuid, total, iva, forma, metodo="PUE", fecha="15/06/2026",
             receptor="AAA010101AAA", clave="01010101"):
    metodo_txt = {"PUE": "PUE - x", "PPD": "PPD - x"}.get(metodo, metodo)
    return {
        "CFDI_Type": "Invoice", "Tipo": "Factura", "UUID": uuid,
        "RFC Emisor": "BBB020202BBB", "Nombre Emisor": "Proveedor",
        "RFC Receptor": receptor, "Metodo de Pago": metodo_txt,
        "FormaDePago": (f"{forma} - x" if forma else None),
        "Fecha Emision": fecha, "SubTotal": total - iva, "IVA 16%": iva,
        "Total": total, "Tipo De Cambio": 1.0,
        "ConceptosDetalle": [{"ClaveProdServ": clave}],
    }


class _ExcelTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.out = os.path.join(self._tmp.name, "reporte.xlsx")
        self.invoices = [
            _factura("U1", 1160.0, 160.0, "03"),                 # deducible
            _factura("U2", 2500.0, 344.8, "01"),                 # no ded efectivo
            _factura("U3", 5800.0, 800.0, None, metodo="PPD"),   # provisionado
        ]

    def tearDown(self):
        self._tmp.cleanup()

    def _export(self, incluir_pt):
        export_to_excel(self.invoices, [], [], self.out, incluir_pt=incluir_pt)
        return openpyxl.load_workbook(self.out)


class TestGate(_ExcelTestCase):
    def test_lite_has_no_pt_sheets(self):
        # Criterio §17.9.7: sin PT, el Excel es el de siempre.
        wb = self._export(incluir_pt=False)
        self.assertIn("Invoices", wb.sheetnames)
        self.assertNotIn(pt_exporter.HOJA_RESUMEN, wb.sheetnames)
        self.assertNotIn(pt_exporter.HOJA_DETALLE, wb.sheetnames)
        self.assertNotIn(pt_exporter.HOJA_FUNDAMENTOS, wb.sheetnames)

    def test_pro_adds_the_three_pt_sheets(self):
        wb = self._export(incluir_pt=True)
        for hoja in (pt_exporter.HOJA_RESUMEN, pt_exporter.HOJA_DETALLE,
                     pt_exporter.HOJA_FUNDAMENTOS):
            self.assertIn(hoja, wb.sheetnames)

    def test_pro_keeps_the_base_invoices_sheet_intact(self):
        # El PT se AGREGA; no toca las hojas base.
        wb = self._export(incluir_pt=True)
        ws = wb["Invoices"]
        self.assertEqual(ws.max_row, 1 + len(self.invoices))  # encabezado + filas


class TestConciliadoBancoVivo(_ExcelTestCase):
    def test_detalle_has_live_sumif_formula(self):
        # Criterio §17.9.6: al marcar la casilla, la suma recalcula en Excel.
        wb = self._export(incluir_pt=True)
        det = wb[pt_exporter.HOJA_DETALLE]
        formulas = [det.cell(r, c).value
                    for r in range(1, det.max_row + 1)
                    for c in range(1, det.max_column + 1)
                    if isinstance(det.cell(r, c).value, str)
                    and det.cell(r, c).value.startswith("=SUMIF")]
        self.assertTrue(formulas, "Se esperaba al menos una fórmula SUMIF viva.")
        # La fórmula suma la columna Total MXN filtrando por la marca ✓.
        self.assertTrue(any("✓" in f for f in formulas))

    def test_detalle_has_conciliado_column_empty_for_user(self):
        wb = self._export(incluir_pt=True)
        det = wb[pt_exporter.HOJA_DETALLE]
        encabezados = [det.cell(1, c).value for c in range(1, det.max_column + 1)]
        self.assertIn("Conciliado banco", encabezados)
        # La columna arranca vacía: la marca la pone el usuario.
        col = encabezados.index("Conciliado banco") + 1
        self.assertIsNone(det.cell(2, col).value)


class TestResumenCuadre(_ExcelTestCase):
    def test_summary_totals_match_source_data(self):
        # Criterio §17.9.1: el resumen cuadra contra los datos de origen.
        wb = self._export(incluir_pt=True)
        res = wb[pt_exporter.HOJA_RESUMEN]
        encabezados = [res.cell(1, c).value for c in range(1, res.max_column + 1)]
        col_total_neto = encabezados.index("Total neto") + 1

        # Suma de "Total neto" de las filas de mes (excluye encabezado y la
        # fila TOTAL final, que es una fórmula).
        suma_meses = 0.0
        for r in range(2, res.max_row):  # sin la fila TOTAL
            val = res.cell(r, col_total_neto).value
            if isinstance(val, (int, float)):
                suma_meses += val
        # Sin notas de crédito, el total neto == suma de Totales de las facturas.
        esperado = sum(f["Total"] for f in self.invoices)
        self.assertAlmostEqual(suma_meses, esperado, places=2)

    def test_status_semaphore_colors_applied(self):
        wb = self._export(incluir_pt=True)
        det = wb[pt_exporter.HOJA_DETALLE]
        encabezados = [det.cell(1, c).value for c in range(1, det.max_column + 1)]
        col_estatus = encabezados.index("Estatus") + 1
        colores = {det.cell(r, col_estatus).fill.fgColor.rgb
                   for r in range(2, 2 + len(self.invoices))}
        # Deben aparecer los tres colores de semáforo (verde/rojo/ámbar).
        self.assertIn("00C6EFCE", colores)  # verde
        self.assertIn("00FFC7CE", colores)  # rojo
        self.assertIn("00FFEB9C", colores)  # ámbar


class TestNoRecibidas(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.out = os.path.join(self._tmp.name, "emitidas.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_pure_emitidas_batch_creates_no_pt_recibidas_sheets(self):
        # El sujeto (SUJ) EMITE a dos clientes distintos: aparece como emisor en
        # TODOS los CFDI, asi que se infiere como sujeto, y ninguna factura lo
        # tiene como receptor -> no hay recibidas -> no se crean hojas PT.
        emitidas = [
            {"CFDI_Type": "Invoice", "Tipo": "Factura", "UUID": "E1",
             "RFC Emisor": "SUJ010101AAA", "RFC Receptor": "CLI010101AAA",
             "Metodo de Pago": "PUE - x", "FormaDePago": "03 - x",
             "Fecha Emision": "15/06/2026", "SubTotal": 1000.0, "IVA 16%": 160.0,
             "Total": 1160.0, "Tipo De Cambio": 1.0,
             "ConceptosDetalle": [{"ClaveProdServ": "01010101"}]},
            {"CFDI_Type": "Invoice", "Tipo": "Factura", "UUID": "E2",
             "RFC Emisor": "SUJ010101AAA", "RFC Receptor": "CLI020202BBB",
             "Metodo de Pago": "PUE - x", "FormaDePago": "03 - x",
             "Fecha Emision": "16/06/2026", "SubTotal": 2000.0, "IVA 16%": 320.0,
             "Total": 2320.0, "Tipo De Cambio": 1.0,
             "ConceptosDetalle": [{"ClaveProdServ": "01010101"}]},
        ]
        export_to_excel(emitidas, [], [], self.out, incluir_pt=True)
        wb = openpyxl.load_workbook(self.out)
        self.assertIn("Invoices", wb.sheetnames)
        self.assertNotIn(pt_exporter.HOJA_RESUMEN, wb.sheetnames)
        self.assertNotIn(pt_exporter.HOJA_DETALLE, wb.sheetnames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
