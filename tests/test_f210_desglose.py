"""
Pruebas de F2.10 etapa 2 (docs/design-f2.10-desglose-conceptos.md, D50-D54):
capa pura reports/desglose_conceptos.py y exporters/desglose_exporter.py.

Fixtures: tests/fixtures/edenred/ (anonimizados) + XML-Test/ para el caso
generico multi-concepto.
"""
import copy
import glob
import os
import sys
import tempfile
import unittest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(REPO_ROOT, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

import openpyxl  # noqa: E402

from cfdi_app.parsers.xml_parser_40 import parse_cfdi_40_invoice  # noqa: E402
from cfdi_app.reports import desglose_conceptos as dc  # noqa: E402
from cfdi_app.exporters import desglose_exporter as dx  # noqa: E402

EDENRED_DIR = os.path.join(REPO_ROOT, "tests", "fixtures", "edenred")
XML_TEST_DIR = os.path.join(REPO_ROOT, "XML-Test")

UUID_JUNIO = "3358E192-7DAB-45A3-AEBE-30EADA074BDC"
UUID_AGOSTO = "46613D69-8126-672B-547F-3F98CD5132A6"
UUID_PEAJE = "29D10416-3337-404A-BF21-8488E8BC522D"
UUID_TAG = "57A56A0B-7D60-490C-B92E-073D80B706C0"


def _cargar(directorio):
    registros = []
    for ruta in sorted(glob.glob(os.path.join(directorio, "*.xml"))):
        data = parse_cfdi_40_invoice(ruta)
        if data:
            registros.append(data)
    return registros


def _por_uuid(registros, uuid):
    return next(r for r in registros if r["UUID"] == uuid)


class _Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.registros = _cargar(EDENRED_DIR)
        cls.lineas = dc.aplanar_lote(cls.registros)


class TestAplanado(_Base):
    def test_junio_trae_comision_mas_33_cargas(self):
        l = dc.lineas_de_uuid(self.lineas, UUID_JUNIO)
        self.assertEqual(sum(1 for x in l if x["origen"] == "CFDI"), 1)
        ecc = [x for x in l if x["origen"] == "ECC"]
        self.assertEqual(len(ecc), 33)
        self.assertAlmostEqual(sum(x["importe"] for x in ecc), 24393.77, 2)
        self.assertAlmostEqual(sum(x["iva"] for x in ecc), 3791.92, 2)
        self.assertAlmostEqual(sum(x["cantidad"] for x in ecc), 1127.97, 2)
        self.assertTrue(all(x["categoria"] == dc.CATEGORIA_ECC for x in ecc))

    def test_la_comision_conserva_su_iva_aparte(self):
        comision = next(x for x in dc.lineas_de_uuid(self.lineas, UUID_JUNIO)
                        if x["origen"] == "CFDI")
        self.assertEqual(comision["descripcion"], "COMISION SOBRE PEDIDOS")
        self.assertAlmostEqual(comision["importe"], 930.0, 2)
        self.assertAlmostEqual(comision["iva"], 148.8, 2)
        self.assertAlmostEqual(comision["total_linea"], 1078.8, 2)

    def test_telepeaje_una_linea_37_cruces(self):
        l = dc.lineas_de_uuid(self.lineas, UUID_PEAJE)
        self.assertEqual(len(l), 1)
        self.assertEqual(l[0]["origen"], "CFDI")
        self.assertAlmostEqual(l[0]["cantidad"], 37.0)
        self.assertAlmostEqual(l[0]["total_linea"], 5772.99, 2)

    def test_etiqueta_tag_total_cero_no_se_filtra(self):
        l = dc.lineas_de_uuid(self.lineas, UUID_TAG)
        self.assertEqual(len(l), 1)
        self.assertAlmostEqual(l[0]["importe"], 1884.0, 2)
        self.assertAlmostEqual(l[0]["descuento"], 1884.0, 2)
        self.assertAlmostEqual(l[0]["total_linea"], 0.0, 2)

    def test_tarjeta_estacion_y_fecha_de_carga(self):
        carga = next(x for x in self.lineas
                     if x["folio_operacion"] == "543884498")
        self.assertEqual(carga["fecha_carga"], "2026-06-24")
        self.assertEqual(carga["rfc_estacion"], "OMC100818GZ7")
        self.assertEqual(carga["clave_estacion"], "11990")
        self.assertEqual(carga["unidad"], "L")
        self.assertTrue(carga["tarjeta"].startswith("9999"))

    def test_mes_ancla_es_el_del_cfdi(self):
        # D51: aunque la carga sea de otro dia, el mes es el de la emision.
        for x in dc.lineas_de_uuid(self.lineas, UUID_AGOSTO):
            self.assertEqual(x["mes_clave"], "2026-08")

    def test_pago_y_nomina_no_entran(self):
        pago = {"CFDI_Type": "Pago", "Tipo": "Pago", "UUID": "X",
                "ConceptosDetalle": [{"Importe": "1"}]}
        nomina = {"CFDI_Type": "Nomina", "Tipo": "Nómina", "UUID": "Y",
                  "ConceptosDetalle": [{"Importe": "1"}]}
        self.assertEqual(dc.aplanar_lote([pago, nomina]), [])


class TestNotaCreditoYMoneda(_Base):
    def _factura(self, tipo, tc=1.0):
        base = copy.deepcopy(_por_uuid(self.registros, UUID_PEAJE))
        base["Tipo"] = tipo
        base["Tipo De Cambio"] = tc
        return base

    def test_nota_de_credito_resta_en_por_proveedor(self):
        factura = self._factura("Factura")
        nota = self._factura("NotaCredito")
        nota["UUID"] = "AAAAAAAA-0000-0000-0000-000000000001"
        nota["Total"] = 1000.0
        # Una NC por 1000 (concepto y total ajustados para que cuadre).
        nota["ConceptosDetalle"] = [{
            "ClaveProdServ": "95111603", "Cantidad": "1", "Unidad": "VIAJE",
            "Descripcion": "DEVOLUCION", "ValorUnitario": "862.07",
            "Importe": "862.07", "Descuento": "",
            "Impuestos": {"Traslados": [{
                "Base": "862.07", "Impuesto": "002", "TipoFactor": "Tasa",
                "TasaOCuota": "0.160000", "Importe": "137.93"}],
                "Retenciones": []}}]
        lineas = dc.aplanar_lote([factura, nota])
        nc = [l for l in lineas if l["tipo"] == "NotaCredito"]
        self.assertEqual(len(nc), 1)
        self.assertAlmostEqual(nc[0]["total_linea"], -1000.0, 2)
        self.assertAlmostEqual(nc[0]["base_iva_16"], -862.07, 2)
        prov = dc.por_proveedor(lineas)[0]
        self.assertAlmostEqual(prov["total"], 5772.99 - 1000.0, 2)
        self.assertEqual(dc.cuadre([nota])[0]["estatus"], dc.ESTATUS_OK)

    def test_moneda_extranjera_se_normaliza_a_mxn(self):
        usd = self._factura("Factura", tc=20.0)
        usd["Moneda"] = "USD"
        l = dc.aplanar_lote([usd])[0]
        self.assertAlmostEqual(l["total_linea"], 5772.99 * 20.0, 2)
        # La cantidad (cruces) no cambia con el tipo de cambio.
        self.assertAlmostEqual(l["cantidad"], 37.0)


class TestAgregaciones(_Base):
    def test_por_proveedor_un_solo_emisor(self):
        filas = dc.por_proveedor(self.lineas)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0]["rfc_emisor"], "ASE930924SS7")
        self.assertEqual(filas[0]["n_cfdi"], 5)

    def test_por_mes_y_por_ejercicio(self):
        meses = {f["mes"]: f for f in dc.por_mes(self.lineas)}
        self.assertEqual(sorted(meses), ["2026-06", "2026-08", "2026-09"])
        self.assertAlmostEqual(meses["2026-06"]["importe_ecc"], 24393.77, 2)
        # Junio: comision 930 + 33 cargas + etiqueta TAG (1,884 con 100% de
        # descuento, mismo mes de emision).
        self.assertAlmostEqual(
            meses["2026-06"]["importe"], 930 + 24393.77 + 1884, 2)
        self.assertAlmostEqual(meses["2026-06"]["descuento"], 1884, 2)
        ejercicios = dc.por_ejercicio(self.lineas)
        self.assertEqual(len(ejercicios), 1)
        self.assertEqual(ejercicios[0]["ejercicio"], 2026)

    def test_iva_del_cfdi_y_del_complemento_nunca_se_mezclan(self):
        # D54: junio = 148.80 (comision) + 3,791.92 (gasolina) = 3,940.72.
        junio = next(f for f in dc.por_iva(self.lineas) if f["mes"] == "2026-06")
        self.assertAlmostEqual(junio["iva_cfdi"], 148.80, 2)
        self.assertAlmostEqual(junio["iva_ecc"], 3791.92, 2)
        self.assertAlmostEqual(junio["iva_total"], 3940.72, 2)
        # Base 16%: comision 930 + importe de las cargas + base de $0.01 que
        # el emisor declara en la etiqueta TAG (viene asi en el XML).
        self.assertAlmostEqual(junio["base_iva_16"], 930 + 24393.77 + 0.01, 2)

    def test_combustible_agrupa_por_tipo_no_por_nombre(self):
        datos = dc.combustible(self.lineas)
        agosto_1 = next(f for f in datos["por_tipo"]
                        if f["mes"] == "2026-08" and f["tipo_combustible"] == "1")
        self.assertEqual(agosto_1["nombres"], "BP REGU, Magna")
        self.assertEqual(agosto_1["cargas"], 31)
        self.assertAlmostEqual(agosto_1["precio_promedio"],
                               agosto_1["importe"] / agosto_1["cantidad"], 2)
        self.assertAlmostEqual(
            sum(f["cantidad"] for f in datos["por_tipo"]), 1127.97 + 1603.79, 2)
        self.assertEqual(
            len({f["tarjeta"] for f in datos["por_tarjeta"]}), 5)

    def test_proveedores_y_filtro(self):
        self.assertEqual(dc.proveedores(self.lineas),
                         [("ASE930924SS7", "EDENRED MEXICO", 5)])
        self.assertEqual(len(dc.filtrar_por_proveedor(self.lineas, "NOEXISTE")), 0)
        self.assertEqual(len(dc.filtrar_por_proveedor(self.lineas, "")),
                         len(self.lineas))

    def test_nombre_de_archivo_derivado_del_lote(self):
        self.assertEqual(dc.nombre_archivo(self.lineas),
                         "Desglose_GAGG010101000_2026-06_a_2026-09.xlsx")
        solo_junio = dc.lineas_de_uuid(self.lineas, UUID_JUNIO)
        self.assertEqual(dc.nombre_archivo(solo_junio, "ASE930924SS7"),
                         "Desglose_GAGG010101000_2026-06_ASE930924SS7.xlsx")


class TestCuadre(_Base):
    def test_edenred_cuadra_al_centavo(self):
        filas = dc.cuadre(self.registros)
        self.assertEqual(len(filas), 5)
        self.assertEqual({f["estatus"] for f in filas}, {dc.ESTATUS_OK})
        self.assertEqual(dc.cuadre_fallas(filas), [])

    def test_cuadre_forzado_a_fallar(self):
        alterado = copy.deepcopy(self.registros)
        junio = _por_uuid(alterado, UUID_JUNIO)
        junio["EstadoCuentaCombustible"]["Conceptos"][0]["Importe"] = "9999.99"
        fallas = dc.cuadre_fallas(dc.cuadre(alterado))
        self.assertEqual([f["uuid"] for f in fallas], [UUID_JUNIO])
        self.assertEqual(fallas[0]["estatus"], dc.ESTATUS_NO_CUADRA)
        self.assertIn("complemento", fallas[0]["detalle"])

    def test_cfdi_sin_detalle_por_concepto_se_avisa(self):
        # Hoy el parser 3.3 no produce ConceptosDetalle (diseño §3.1).
        viejo = {"CFDI_Type": "Invoice", "Tipo": "Factura", "Version": "3.3",
                 "UUID": "BBBBBBBB-0000-0000-0000-000000000001",
                 "Fecha Emision": "01/02/2019", "Total": 116.0}
        self.assertEqual(dc.aplanar_lote([viejo]), [])
        fila = dc.cuadre([viejo])[0]
        self.assertEqual(fila["estatus"], dc.ESTATUS_SIN_DETALLE)

    def test_todos_los_fixtures_de_xml_test_cuadran(self):
        # Incluye el CFDI de hospedaje con ISH (impuesto local): su Total
        # trae $80.69 que no cuelga de ningun concepto -> linea explicita.
        registros = _cargar(XML_TEST_DIR)
        self.assertGreater(len(registros), 0)
        self.assertEqual(dc.cuadre_fallas(dc.cuadre(registros)), [])

    def test_impuestos_locales_son_una_linea_explicita(self):
        registros = _cargar(XML_TEST_DIR)
        hotel = _por_uuid(registros, "9A655B6C-20C8-4A20-942E-F50D77D45201")
        lineas = dc.aplanar_lote([hotel])
        locales = [l for l in lineas
                   if l["categoria"] == dc.CATEGORIA_IMPUESTOS_LOCALES]
        self.assertEqual(len(locales), 1)
        self.assertAlmostEqual(locales[0]["total_linea"], 80.69, 2)
        self.assertAlmostEqual(sum(l["total_linea"] for l in lineas), 3200.63, 2)


class TestReportePorXml(_Base):
    """Etapa 3: helpers puros del dialogo "Ver desglose de este XML"."""

    def test_tabla_de_estado_de_combustible_incluye_columnas_de_carga(self):
        l = dc.lineas_de_uuid(self.lineas, UUID_JUNIO)
        encabezados, filas, numericas = dc.tabla_xml(l)
        self.assertEqual(len(filas), 34)  # comision + 33 cargas
        self.assertIn("Tarjeta", encabezados)
        self.assertEqual(len(encabezados), len(numericas))
        self.assertTrue(all(len(f) == len(encabezados) for f in filas))
        comision = dict(zip(encabezados, filas[0]))
        self.assertEqual(comision["Origen"], "CFDI")
        self.assertEqual(comision["Total"], "1,078.80")
        self.assertEqual(comision["Tarjeta"], "")
        carga = dict(zip(encabezados, filas[1]))
        self.assertEqual(carga["Origen"], "ECC")
        self.assertEqual(carga["Descripción"], "Magna")
        self.assertEqual(carga["Fecha carga"], "2026-06-24")

    def test_factura_normal_no_muestra_columnas_de_combustible(self):
        encabezados, filas, _ = dc.tabla_xml(
            dc.lineas_de_uuid(self.lineas, UUID_PEAJE))
        self.assertEqual(len(filas), 1)
        self.assertNotIn("Tarjeta", encabezados)
        self.assertNotIn("Fecha carga", encabezados)

    def test_resumen_separa_iva_del_cfdi_y_del_complemento(self):
        r = dc.resumen_xml(dc.lineas_de_uuid(self.lineas, UUID_JUNIO))
        self.assertEqual(r["lineas_cfdi"], 1)
        self.assertEqual(r["cargas"], 33)
        self.assertAlmostEqual(r["importe_cfdi"], 930.0, 2)
        self.assertAlmostEqual(r["importe_ecc"], 24393.77, 2)
        self.assertAlmostEqual(r["iva_cfdi"], 148.80, 2)
        self.assertAlmostEqual(r["iva_ecc"], 3791.92, 2)
        self.assertAlmostEqual(r["iva_total"], 3940.72, 2)
        self.assertAlmostEqual(r["total_cfdi"], 1078.80, 2)
        self.assertAlmostEqual(r["total_ecc"], 28185.69, 2)
        self.assertAlmostEqual(r["litros"], 1127.97, 2)

    def test_pie_sin_complemento_es_un_solo_renglon(self):
        pie = dc.texto_pie_xml(
            dc.resumen_xml(dc.lineas_de_uuid(self.lineas, UUID_PEAJE)))
        self.assertEqual(len(pie), 1)
        self.assertIn("$5,772.99", pie[0])

    def test_pie_con_complemento_trae_el_iva_total(self):
        pie = dc.texto_pie_xml(
            dc.resumen_xml(dc.lineas_de_uuid(self.lineas, UUID_JUNIO)))
        self.assertEqual(len(pie), 3)
        self.assertIn("IVA total (CFDI + complemento): $3,940.72", pie[2])

    def test_nombre_de_archivo_de_un_xml_lleva_la_factura_saneada(self):
        l = dc.lineas_de_uuid(self.lineas, UUID_JUNIO)
        self.assertEqual(
            dc.nombre_archivo(l, sufijo="TCGE-000004231229"),
            "Desglose_GAGG010101000_2026-06_TCGE-000004231229.xlsx")
        # Caracteres no validos en un nombre de archivo se neutralizan.
        self.assertEqual(
            dc.nombre_archivo(l, sufijo='A/B:C*"1'),
            "Desglose_GAGG010101000_2026-06_A-B-C-1.xlsx")

    def test_exportar_un_solo_xml_con_las_hojas_de_la_etapa_3(self):
        junio = [_por_uuid(self.registros, UUID_JUNIO)]
        ruta = os.path.join(tempfile.mkdtemp(), "uno.xlsx")
        hojas = {dx.HOJA_DETALLE, dx.HOJA_IVA, dx.HOJA_CUADRE,
                 dx.HOJA_COMBUSTIBLE}
        escritas = dx.exportar_desglose(junio, ruta, hojas=hojas)
        self.assertEqual(escritas, [dx.HOJA_DETALLE, dx.HOJA_COMBUSTIBLE,
                                    dx.HOJA_IVA, dx.HOJA_CUADRE])
        ws = openpyxl.load_workbook(ruta)[dx.HOJA_DETALLE]
        self.assertEqual(ws.max_row, 1 + 34 + 2)


class TestFiltroDeProveedores(unittest.TestCase):
    PROVS = [("ASE930924SS7", "EDENRED MEXICO", 5),
             ("WMT970714R10", "NUEVA WAL MART DE MÉXICO", 12),
             ("HDM001017AS1", "HOME DEPOT MEXICO", 3)]

    def test_vacio_devuelve_todos(self):
        self.assertEqual(dc.filtrar_proveedores(self.PROVS, ""), self.PROVS)
        self.assertEqual(dc.filtrar_proveedores(self.PROVS, "   "), self.PROVS)

    def test_por_rfc_completo_o_parcial_sin_distinguir_mayusculas(self):
        self.assertEqual([p[0] for p in dc.filtrar_proveedores(self.PROVS, "ase9309")],
                         ["ASE930924SS7"])
        self.assertEqual([p[0] for p in dc.filtrar_proveedores(self.PROVS, "r10")],
                         ["WMT970714R10"])

    def test_por_razon_social_parcial(self):
        self.assertEqual([p[0] for p in dc.filtrar_proveedores(self.PROVS, "depot")],
                         ["HDM001017AS1"])

    def test_sin_acentos(self):
        # 'mexico' encuentra 'MÉXICO' y 'MEXICO'.
        self.assertEqual(len(dc.filtrar_proveedores(self.PROVS, "mexico")), 3)
        self.assertEqual([p[0] for p in dc.filtrar_proveedores(self.PROVS, "méxico wal")],
                         ["WMT970714R10"])

    def test_varias_palabras_en_cualquier_orden(self):
        self.assertEqual([p[0] for p in dc.filtrar_proveedores(self.PROVS, "mart wal")],
                         ["WMT970714R10"])

    def test_sin_coincidencias(self):
        self.assertEqual(dc.filtrar_proveedores(self.PROVS, "zzz"), [])


class TestVisorDelLote(_Base):
    """Filtro + arbol Proveedor -> Factura -> lineas del visor del lote."""

    def test_filtro_vacio_devuelve_todo(self):
        self.assertEqual(len(dc.filtrar_lineas(self.lineas, "")), len(self.lineas))
        self.assertEqual(len(dc.filtrar_lineas(self.lineas, "   ")), len(self.lineas))

    def test_busca_por_descripcion_sin_distinguir_mayusculas(self):
        magna = dc.filtrar_lineas(self.lineas, "MAGNA")
        self.assertGreater(len(magna), 0)
        self.assertTrue(all(l["descripcion"] == "Magna" for l in magna))

    def test_busca_por_rfc_de_la_estacion_y_por_factura(self):
        por_estacion = dc.filtrar_lineas(self.lineas, "omc100818gz7")
        self.assertEqual(len(por_estacion), 1)
        self.assertEqual(por_estacion[0]["factura"], "TCGE-000004231229")
        por_factura = dc.filtrar_lineas(self.lineas, "TCTL-000000080149")
        self.assertEqual({l["uuid"] for l in por_factura}, {UUID_PEAJE})

    def test_busca_por_uuid_y_por_proveedor(self):
        self.assertEqual(
            {l["uuid"] for l in dc.filtrar_lineas(self.lineas, "3358e192")},
            {UUID_JUNIO})
        self.assertEqual(len(dc.filtrar_lineas(self.lineas, "edenred")),
                         len(self.lineas))

    def test_montos_con_la_misma_sintaxis_del_visor(self):
        grandes = dc.filtrar_lineas(self.lineas, "total>=5000")
        self.assertEqual([l["uuid"] for l in grandes], [UUID_PEAJE])
        # Sin nombre de campo se usa el total de la linea.
        self.assertEqual(len(dc.filtrar_lineas(self.lineas, ">=5000")), 1)
        iva = dc.filtrar_lineas(self.lineas, "iva>=300")
        self.assertTrue(all(l["iva"] >= 300 for l in iva))

    def test_sin_coincidencias(self):
        self.assertEqual(dc.filtrar_lineas(self.lineas, "tornillo"), [])
        self.assertEqual(dc.arbol([]), [])

    def test_arbol_proveedor_factura_lineas(self):
        nodos = dc.arbol(self.lineas)
        self.assertEqual(len(nodos), 1)
        prov = nodos[0]
        self.assertTrue(prov["texto"].startswith("ASE930924SS7 – EDENRED MEXICO"))
        self.assertIsNone(prov["uuid"])
        self.assertEqual(len(prov["hijos"]), 5)  # 5 facturas
        self.assertEqual(len(dc.COLUMNAS_ARBOL), len(prov["cols"]))
        junio = next(f for f in prov["hijos"] if f["uuid"] == UUID_JUNIO)
        self.assertTrue(junio["texto"].startswith("TCGE-000004231229"))
        self.assertIn(UUID_JUNIO, junio["texto"])
        self.assertEqual(len(junio["hijos"]), 34)
        self.assertEqual(junio["cols"][0], "2026-06-30")
        # Montos de la factura = suma de sus lineas (comision + 33 cargas).
        self.assertEqual(junio["cols"][-1], "29,264.49")  # 1,078.80 + 28,185.69
        self.assertTrue(all(h["uuid"] == UUID_JUNIO for h in junio["hijos"]))

    def test_el_arbol_de_lo_filtrado_solo_trae_lo_filtrado(self):
        nodos = dc.arbol(dc.filtrar_lineas(self.lineas, "omc100818gz7"))
        self.assertEqual(len(nodos[0]["hijos"]), 1)
        self.assertEqual(len(nodos[0]["hijos"][0]["hijos"]), 1)
        linea = nodos[0]["hijos"][0]["hijos"][0]
        self.assertIn("estación OMC100818GZ7", linea["texto"])
        self.assertIn("2026-06-24", linea["texto"])
        self.assertEqual(linea["cols"][1], "ECC")

    def test_resumen_del_lote_separa_los_iva(self):
        magna = dc.filtrar_lineas(self.lineas, "magna")
        texto = dc.texto_resumen_lote(magna, len(self.lineas))
        self.assertIn(f"{len(magna)} de {len(self.lineas)} líneas", texto)
        self.assertIn("IVA del CFDI: $0.00", texto)
        self.assertIn("IVA del complemento:", texto)

    def test_exportar_lo_filtrado_no_incluye_cuadre(self):
        magna = dc.filtrar_lineas(self.lineas, "magna")
        ruta = os.path.join(tempfile.mkdtemp(), "filtrado.xlsx")
        escritas = dx.exportar_desglose_lineas(magna, ruta)
        self.assertNotIn(dx.HOJA_CUADRE, escritas)
        self.assertIn(dx.HOJA_DETALLE, escritas)
        ws = openpyxl.load_workbook(ruta)[dx.HOJA_DETALLE]
        self.assertEqual(ws.max_row, 1 + len(magna) + 2)

    def test_nombre_del_archivo_filtrado(self):
        magna = dc.filtrar_lineas(self.lineas, "magna")
        self.assertEqual(
            dc.nombre_archivo(magna, sufijo="filtrado-magna"),
            "Desglose_GAGG010101000_2026-06_a_2026-08_filtrado-magna.xlsx")


class TestExporter(_Base):
    def _exportar(self, **kwargs):
        carpeta = tempfile.mkdtemp()
        ruta = os.path.join(carpeta, "desglose.xlsx")
        escritas = dx.exportar_desglose(self.registros, ruta, **kwargs)
        return escritas, openpyxl.load_workbook(ruta)

    def test_todas_las_hojas_con_ecc(self):
        escritas, wb = self._exportar()
        self.assertEqual(escritas, dx.HOJAS_DISPONIBLES)
        self.assertEqual(wb.sheetnames, dx.HOJAS_DISPONIBLES)

    def test_detalle_una_fila_por_linea_y_subtotal_vivo(self):
        _, wb = self._exportar()
        ws = wb[dx.HOJA_DETALLE]
        self.assertEqual(ws.max_row, 1 + len(self.lineas) + 2)
        self.assertEqual(ws.auto_filter.ref[:2], "A1")
        self.assertEqual(ws.freeze_panes, "A2")
        fila_total = ws.max_row
        formulas = [c.value for c in ws[fila_total] if isinstance(c.value, str)
                    and c.value.startswith("=SUBTOTAL(9,")]
        self.assertGreater(len(formulas), 3)

    def test_hoja_iva_separa_cfdi_y_complemento(self):
        _, wb = self._exportar()
        ws = wb[dx.HOJA_IVA]
        encabezados = [c.value for c in ws[3]]
        self.assertIn("IVA del CFDI", encabezados)
        self.assertIn("IVA del complemento (combustible monedero)", encabezados)
        fila_junio = next(r for r in ws.iter_rows(min_row=4, values_only=True)
                          if r[0] == "2026-06")
        self.assertAlmostEqual(fila_junio[encabezados.index("IVA del CFDI")], 148.8, 2)
        self.assertAlmostEqual(
            fila_junio[encabezados.index("IVA del complemento (combustible monedero)")],
            3791.92, 2)

    def test_hoja_cuadre_dice_que_todos_cuadran(self):
        _, wb = self._exportar()
        self.assertTrue(wb[dx.HOJA_CUADRE]["A1"].value.startswith("Todos cuadran"))

    def test_hoja_cuadre_lista_los_que_fallan(self):
        alterado = copy.deepcopy(self.registros)
        _por_uuid(alterado, UUID_JUNIO)["EstadoCuentaCombustible"][
            "Conceptos"][0]["Importe"] = "9999.99"
        ruta = os.path.join(tempfile.mkdtemp(), "d.xlsx")
        dx.exportar_desglose(alterado, ruta)
        ws = openpyxl.load_workbook(ruta)[dx.HOJA_CUADRE]
        self.assertEqual(ws["C2"].value, UUID_JUNIO)

    def test_seleccion_de_hojas(self):
        escritas, wb = self._exportar(hojas={dx.HOJA_IVA, dx.HOJA_DETALLE})
        self.assertEqual(escritas, [dx.HOJA_DETALLE, dx.HOJA_IVA])
        self.assertEqual(wb.sheetnames, [dx.HOJA_DETALLE, dx.HOJA_IVA])

    def test_combustible_no_se_escribe_sin_ecc(self):
        sin_ecc = [r for r in self.registros if r["UUID"] in (UUID_PEAJE, UUID_TAG)]
        ruta = os.path.join(tempfile.mkdtemp(), "d.xlsx")
        escritas = dx.exportar_desglose(sin_ecc, ruta)
        self.assertNotIn(dx.HOJA_COMBUSTIBLE, escritas)

    def test_filtro_por_proveedor_sin_coincidencias_no_truena(self):
        ruta = os.path.join(tempfile.mkdtemp(), "d.xlsx")
        escritas = dx.exportar_desglose(self.registros, ruta, rfc_emisor="NOEXISTE")
        self.assertIn(dx.HOJA_DETALLE, escritas)

    def test_sin_hojas_elegidas_guarda_un_libro_valido(self):
        ruta = os.path.join(tempfile.mkdtemp(), "d.xlsx")
        dx.exportar_desglose(self.registros, ruta, hojas=set())
        self.assertEqual(openpyxl.load_workbook(ruta).sheetnames, ["Desglose"])


if __name__ == "__main__":
    unittest.main()
