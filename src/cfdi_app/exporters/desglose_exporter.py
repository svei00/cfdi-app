# --- cfdi_app/exporters/desglose_exporter.py ---
"""Excel del reporte "Desglose de conceptos" (F2.10, D52).

Archivo APARTE del Excel principal (D11: el principal y sus hojas PT alimentan
las formulas INDEX/MATCH del dueño; miles de filas de detalle ahi lo romperian).
Consume la capa pura reports/desglose_conceptos.py; aqui solo vive openpyxl.

Cada hoja lleva autofiltro, encabezado congelado y una fila de totales con
=SUBTOTAL(9, ...): al filtrar, suma solo lo visible.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cfdi_app.reports import desglose_conceptos as dc

HOJA_DETALLE = "Detalle"
HOJA_PROVEEDOR = "Por proveedor"
HOJA_MES = "Por mes"
HOJA_EJERCICIO = "Por ejercicio"
HOJA_CONCEPTO = "Por concepto"
HOJA_COMBUSTIBLE = "Combustible"
HOJA_IVA = "IVA"
HOJA_CUADRE = "Cuadre"

# Orden de aparicion en el libro. Combustible solo existe si hay lineas ECC.
HOJAS_DISPONIBLES = [
    HOJA_DETALLE, HOJA_PROVEEDOR, HOJA_MES, HOJA_EJERCICIO, HOJA_CONCEPTO,
    HOJA_COMBUSTIBLE, HOJA_IVA, HOJA_CUADRE,
]

_FMT_DINERO = '#,##0.00'
_FMT_CANTIDAD = '#,##0.00##'
_FMT_PRECIO = '#,##0.000'

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_TOTAL_FONT = Font(bold=True)
_NOTA_FONT = Font(italic=True, color="7F7F7F")
_MAL_FILL = PatternFill("solid", fgColor="FFC7CE")
_AVISO_FILL = PatternFill("solid", fgColor="FFEB9C")

# (clave, encabezado, formato, sumar_en_total). formato None = texto.
_COLS_DETALLE = [
    ("mes_clave", "Mes", None, False),
    ("fecha_cfdi", "Fecha CFDI", None, False),
    ("factura", "Factura", None, False),
    ("uuid", "UUID", None, False),
    ("rfc_emisor", "RFC emisor", None, False),
    ("nombre_emisor", "Proveedor", None, False),
    ("tipo", "Tipo", None, False),
    ("forma_pago", "Forma de pago", None, False),
    ("metodo_pago", "Método de pago", None, False),
    ("origen", "Origen", None, False),
    ("num_linea", "Línea", '0', False),
    ("clave_prod_serv", "ClaveProdServ", None, False),
    ("descripcion", "Descripción", None, False),
    ("categoria", "Categoría", None, False),
    ("cantidad", "Cantidad", _FMT_CANTIDAD, False),  # unidades mixtas: sin total
    ("unidad", "Unidad", None, False),
    ("valor_unitario", "Valor unitario", _FMT_PRECIO, False),
    ("importe", "Importe", _FMT_DINERO, True),
    ("descuento", "Descuento", _FMT_DINERO, True),
    ("iva", "IVA", _FMT_DINERO, True),
    ("ieps", "IEPS", _FMT_DINERO, True),
    ("retenciones", "Retenciones", _FMT_DINERO, True),
    ("total_linea", "Total línea", _FMT_DINERO, True),
    ("tarjeta", "Tarjeta", None, False),
    ("rfc_estacion", "RFC estación", None, False),
    ("clave_estacion", "Clave estación", None, False),
    ("tipo_combustible", "Tipo combustible", None, False),
    ("folio_operacion", "Folio operación", None, False),
    ("fecha_carga", "Fecha carga", None, False),
]

_COLS_METRICAS = [
    ("n_cfdi", "CFDI", '0', True),
    ("n_lineas", "Líneas", '0', True),
    ("importe", "Importe", _FMT_DINERO, True),
    ("descuento", "Descuento", _FMT_DINERO, True),
    ("iva", "IVA", _FMT_DINERO, True),
    ("total", "Total", _FMT_DINERO, True),
]
# Subconjunto del complemento de combustible (D54: visible, nunca mezclado).
_COLS_ECC = [
    ("importe_ecc", "De los cuales ECC: importe", _FMT_DINERO, True),
    ("iva_ecc", "De los cuales ECC: IVA", _FMT_DINERO, True),
    ("total_ecc", "De los cuales ECC: total", _FMT_DINERO, True),
]


def _prov():
    return [("rfc_emisor", "RFC emisor", None, False),
            ("nombre_emisor", "Proveedor", None, False)]


def _cols_resumen(prefijo, con_ecc):
    cols = prefijo + _COLS_METRICAS
    return cols + (_COLS_ECC if con_ecc else [])


def _escribir_tabla(ws, filas, columnas, nota=None, fila_inicial=1,
                    hoja_simple=True):
    """Escribe encabezado, filas y fila SUBTOTAL. Con hoja_simple (una tabla
    por hoja) ademas congela el encabezado, pone autofiltro y ajusta anchos;
    la hoja Combustible apila 3 tablas y desactiva eso (Excel permite UN solo
    autofiltro por hoja). Devuelve la siguiente fila libre."""
    if nota:
        ws.cell(row=fila_inicial, column=1, value=nota).font = _NOTA_FONT
        fila_inicial += 2  # deja una fila en blanco entre nota y tabla

    for j, (_, titulo, _, _) in enumerate(columnas, start=1):
        c = ws.cell(row=fila_inicial, column=j, value=titulo)
        c.font, c.fill = _HEADER_FONT, _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, fila in enumerate(filas, start=fila_inicial + 1):
        for j, (clave, _, fmt, _) in enumerate(columnas, start=1):
            valor = fila.get(clave)
            c = ws.cell(row=i, column=j, value=valor)
            if fmt:
                c.number_format = fmt

    ultima = fila_inicial + len(filas)
    fila_total = ultima + 2
    if filas:
        if hoja_simple:
            ws.auto_filter.ref = (
                f"A{fila_inicial}:{get_column_letter(len(columnas))}{ultima}")
        ws.cell(row=fila_total, column=1, value="TOTAL (filas visibles)").font = _TOTAL_FONT
        for j, (_, _, fmt, sumar) in enumerate(columnas, start=1):
            if not sumar:
                continue
            letra = get_column_letter(j)
            c = ws.cell(row=fila_total, column=j,
                        value=f"=SUBTOTAL(9,{letra}{fila_inicial + 1}:{letra}{ultima})")
            c.font, c.number_format = _TOTAL_FONT, fmt
    if hoja_simple:
        ws.freeze_panes = ws.cell(row=fila_inicial + 1, column=1)
        _ajustar_anchos(ws, columnas)
    return (fila_total if filas else ultima) + 2


def _ajustar_anchos(ws, columnas):
    for j, (_, titulo, _, _) in enumerate(columnas, start=1):
        largo = max(len(titulo), 10)
        for fila in ws.iter_rows(min_row=1, min_col=j, max_col=j, max_row=min(ws.max_row, 200)):
            valor = fila[0].value
            if isinstance(valor, str) and not valor.startswith("="):
                largo = max(largo, min(len(valor), 60))
        ws.column_dimensions[get_column_letter(j)].width = min(largo + 2, 62)


def _hoja_detalle(wb, lineas):
    ws = wb.create_sheet(HOJA_DETALLE)
    _escribir_tabla(ws, lineas, _COLS_DETALLE)


def _hoja_resumen(wb, nombre, filas, prefijo, con_ecc):
    ws = wb.create_sheet(nombre)
    _escribir_tabla(ws, filas, _cols_resumen(prefijo, con_ecc))


def _hoja_concepto(wb, lineas, con_ecc):
    ws = wb.create_sheet(HOJA_CONCEPTO)
    prefijo = _prov() + [("clave_prod_serv", "ClaveProdServ", None, False),
                         ("descripcion", "Descripción", None, False),
                         ("cantidad", "Cantidad", _FMT_CANTIDAD, False)]
    _escribir_tabla(ws, dc.por_concepto(lineas), _cols_resumen(prefijo, con_ecc))


def _hoja_iva(wb, lineas):
    """D54: IVA del CFDI y del complemento SIEMPRE en columnas separadas."""
    ws = wb.create_sheet(HOJA_IVA)
    cols = [("mes", "Mes", None, False)] + _prov() + [
        ("base_iva_16", "Base IVA 16%", _FMT_DINERO, True),
        ("base_iva_8", "Base IVA 8%", _FMT_DINERO, True),
        ("base_iva_0", "Base IVA 0%", _FMT_DINERO, True),
        ("base_iva_exenta", "Base exenta", _FMT_DINERO, True),
        ("iva_cfdi", "IVA del CFDI", _FMT_DINERO, True),
        ("iva_ecc", "IVA del complemento (combustible monedero)", _FMT_DINERO, True),
        ("iva_total", "IVA total", _FMT_DINERO, True),
    ]
    nota = ("El IVA del CFDI y el del complemento de combustible se muestran "
            "separados a propósito: el Total de un estado de cuenta de "
            "monedero (p. ej. Edenred) es solo la comisión; el combustible "
            "vive en el complemento.")
    _escribir_tabla(ws, dc.por_iva(lineas), cols, nota=nota)


def _hoja_combustible(wb, lineas):
    ws = wb.create_sheet(HOJA_COMBUSTIBLE)
    datos = dc.combustible(lineas)
    comunes = [("cargas", "Cargas", '0', True),
               ("cantidad", "Litros", _FMT_CANTIDAD, True),
               ("importe", "Importe", _FMT_DINERO, True),
               ("iva", "IVA", _FMT_DINERO, True),
               ("total", "Total", _FMT_DINERO, True),
               ("precio_promedio", "Precio promedio por litro", _FMT_PRECIO, False)]
    bloques = [
        ("Por mes y tipo de combustible (el código se muestra tal cual; el "
         "nombre depende de la marca)",
         datos["por_tipo"],
         [("mes", "Mes", None, False),
          ("tipo_combustible", "Tipo combustible", None, False),
          ("nombres", "Nombres vistos", None, False)] + comunes),
        ("Por mes y tarjeta",
         datos["por_tarjeta"],
         [("mes", "Mes", None, False), ("tarjeta", "Tarjeta", None, False)] + comunes),
        ("Por estación",
         datos["por_estacion"],
         [("rfc_estacion", "RFC estación", None, False),
          ("clave_estacion", "Clave estación", None, False)] + comunes),
    ]
    fila = 1
    for titulo, filas, cols in bloques:
        ws.cell(row=fila, column=1, value=titulo).font = _TOTAL_FONT
        fila = _escribir_tabla(ws, filas, cols, fila_inicial=fila + 1,
                               hoja_simple=False)
    for j in range(1, 12):
        ws.column_dimensions[get_column_letter(j)].width = 22


def _hoja_cuadre(wb, filas_cuadre):
    ws = wb.create_sheet(HOJA_CUADRE)
    fallas = dc.cuadre_fallas(filas_cuadre)
    if not fallas:
        ws.cell(row=1, column=1,
                value=f"Todos cuadran ({len(filas_cuadre)} CFDI revisados, "
                      f"tolerancia ±{dc.TOLERANCIA_CUADRE:.2f}).").font = _TOTAL_FONT
        ws.column_dimensions["A"].width = 70
        return
    cols = [("mes", "Mes", None, False), ("factura", "Factura", None, False),
            ("uuid", "UUID", None, False)] + _prov() + [
        ("version", "Versión CFDI", None, False),
        ("suma_lineas", "Σ líneas", _FMT_DINERO, False),
        ("total_declarado", "Total CFDI", _FMT_DINERO, False),
        ("ecc_importe", "Σ cargas (importe)", _FMT_DINERO, False),
        ("ecc_subtotal_declarado", "SubTotal complemento", _FMT_DINERO, False),
        ("ecc_total", "Σ cargas + IVA", _FMT_DINERO, False),
        ("ecc_total_declarado", "Total complemento", _FMT_DINERO, False),
        ("estatus", "Estatus", None, False),
        ("detalle", "Detalle", None, False),
    ]
    _escribir_tabla(ws, fallas, cols)
    col_estatus = [c[0] for c in cols].index("estatus") + 1
    for i in range(2, 2 + len(fallas)):
        celda = ws.cell(row=i, column=col_estatus)
        celda.fill = _MAL_FILL if celda.value == dc.ESTATUS_NO_CUADRA else _AVISO_FILL


def exportar_desglose(parsed_data_list, output_path, rfc_emisor=None, hojas=None):
    """Genera el Excel de Desglose de conceptos.

    parsed_data_list: lista de dicts del parser (ProcessResult.all_parsed_data).
    rfc_emisor: None/'' = todos los proveedores; un RFC = solo ese proveedor.
    hojas: conjunto de nombres de HOJAS_DISPONIBLES a incluir (default: todas).
           'Combustible' solo se escribe si el lote tiene cargas ECC.
    Devuelve la lista de hojas escritas. Puede propagar PermissionError (archivo
    abierto en Excel); el llamador decide como mostrarlo.
    """
    elegidas = set(HOJAS_DISPONIBLES if hojas is None else hojas)
    todas = dc.aplanar_lote(parsed_data_list)
    lineas = dc.filtrar_por_proveedor(todas, rfc_emisor)
    cuadre_filas = dc.cuadre(
        [d for d in parsed_data_list
         if not rfc_emisor or d.get("RFC Emisor") == rfc_emisor])
    return _escribir_libro(lineas, cuadre_filas, output_path, elegidas)


def exportar_desglose_lineas(lineas, output_path, hojas=None):
    """Excel del desglose a partir de LINEAS ya filtradas (visor del lote).
    La hoja Cuadre NUNCA se incluye: compara contra el Total del CFDI
    completo y, con solo algunas lineas, saldria "no cuadra" falsamente."""
    elegidas = set(HOJAS_DISPONIBLES if hojas is None else hojas) - {HOJA_CUADRE}
    return _escribir_libro(lineas, [], output_path, elegidas)


def _escribir_libro(lineas, cuadre_filas, output_path, elegidas):
    con_ecc = dc.hay_ecc(lineas)

    wb = Workbook()
    wb.remove(wb.active)
    escritas = []

    def agregar(nombre, escribir):
        if nombre in elegidas:
            escribir()
            escritas.append(nombre)

    agregar(HOJA_DETALLE, lambda: _hoja_detalle(wb, lineas))
    agregar(HOJA_PROVEEDOR, lambda: _hoja_resumen(
        wb, HOJA_PROVEEDOR, dc.por_proveedor(lineas), _prov(), con_ecc))
    agregar(HOJA_MES, lambda: _hoja_resumen(
        wb, HOJA_MES, dc.por_mes(lineas),
        [("mes", "Mes", None, False)] + _prov(), con_ecc))
    agregar(HOJA_EJERCICIO, lambda: _hoja_resumen(
        wb, HOJA_EJERCICIO, dc.por_ejercicio(lineas),
        [("ejercicio", "Ejercicio", None, False)] + _prov(), con_ecc))
    agregar(HOJA_CONCEPTO, lambda: _hoja_concepto(wb, lineas, con_ecc))
    if con_ecc:
        agregar(HOJA_COMBUSTIBLE, lambda: _hoja_combustible(wb, lineas))
    agregar(HOJA_IVA, lambda: _hoja_iva(wb, lineas))
    agregar(HOJA_CUADRE, lambda: _hoja_cuadre(wb, cuadre_filas))

    if not escritas:  # openpyxl no guarda un libro sin hojas
        ws = wb.create_sheet("Desglose")
        ws.cell(row=1, column=1, value="No se eligió ninguna hoja.")
    wb.save(output_path)
    return escritas
